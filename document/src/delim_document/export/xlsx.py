"""Readable XLSX report renderer."""

from __future__ import annotations

from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from delim_document.export.csv import HEADERS
from delim_document.export.models import ReportRow


def _excel_datetime(value: datetime) -> datetime:
    """Return the same instant as a timezone-naive UTC datetime for Excel."""
    if value.tzinfo is None:
        return value
    return value.astimezone(UTC).replace(tzinfo=None)


def render_xlsx(group_name: str, rows: tuple[ReportRow, ...]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Операции"
    sheet.append(("Группа", group_name))
    sheet.append(())
    sheet.append(HEADERS)
    for cell in sheet[3]:
        cell.font = Font(bold=True)

    for row in rows:
        sheet.append(
            (
                _excel_datetime(row.date),
                row.description,
                row.payer,
                Decimal(row.amount_minor) / Decimal(100),
                row.currency,
                row.note,
            )
        )
        sheet.cell(sheet.max_row, 1).number_format = "yyyy-mm-dd hh:mm:ss"
        sheet.cell(sheet.max_row, 4).number_format = "0.00"

    for column in range(1, len(HEADERS) + 1):
        values = (
            "" if cell.value is None else str(cell.value)
            for cell in sheet[get_column_letter(column)]
        )
        width = min(45, max((len(value) for value in values), default=0) + 2)
        sheet.column_dimensions[get_column_letter(column)].width = max(10, width)
    sheet.freeze_panes = "A4"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
